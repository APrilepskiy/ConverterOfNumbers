"""Веб-приложение для конвертации файлов Apple Numbers (.numbers) в Excel (.xlsx)."""

import io
import os
import random
import string
import tempfile
from datetime import datetime

from flask import Flask, jsonify, render_template, request, send_file
from numbers_parser import Document

app = Flask(__name__)

MAX_FILE_SIZE = 100 * 1024 * 1024  # 100 МБ
ALLOWED_EXTENSIONS = {".numbers"}


def allowed_file(filename: str) -> bool:
    return os.path.splitext(filename.lower())[1] in ALLOWED_EXTENSIONS


@app.route("/")
def index():
    return render_template("index.html")


@app.route("/api/convert", methods=["POST"])
def convert():
    file = request.files.get("file")
    if file is None or file.filename == "":
        return jsonify({"error": "Файл не выбран."}), 400

    filename = file.filename
    if not allowed_file(filename):
        return jsonify({"error": "Поддерживаются только файлы формата .numbers."}), 400

    # Проверка размера файла
    file.seek(0, io.SEEK_END)
    size = file.tell()
    file.seek(0)
    if size > MAX_FILE_SIZE:
        return jsonify({"error": "Файл слишком большой (максимум 100 МБ)."}), 413

    try:
        with tempfile.NamedTemporaryFile(suffix=".numbers", delete=True) as tmp:
            file.save(tmp)
            tmp.flush()
            doc = Document(filename=tmp.name)
    except Exception:
        return jsonify({"error": "Не удалось прочитать файл. "
                                 "Убедитесь, что это корректный файл Apple Numbers."}), 400

    try:
        xlsx_stream = convert_to_xlsx(doc)
    except Exception as exc:
        return jsonify({"error": f"Ошибка при конвертации: {exc}"}), 500

    base = os.path.splitext(os.path.basename(filename))[0]
    out_name = f"{base}.xlsx"
    timestamp = datetime.now().strftime("%Y%m%d%H%M%S")
    suffix = "".join(random.choices(string.ascii_lowercase + string.digits, k=6))
    download_name = f"{timestamp}_{suffix}_{out_name}"

    return send_file(
        xlsx_stream,
        as_attachment=True,
        download_name=download_name,
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )


def convert_to_xlsx(doc: Document) -> io.BytesIO:
    """Конвертирует документ Numbers в xlsx, сохраняя структуру листов и данные ячеек."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    wb.remove(wb.active)

    for sheet in doc.sheets:
        for table in sheet.tables:
            ws = wb.create_sheet(title=sanitize_sheet_name(table.name))

            num_rows = table.num_rows
            num_cols = table.num_cols

            for r in range(num_rows):
                for c in range(num_cols):
                    cell = table.cell(r, c)
                    value = cell.value
                    if value is None:
                        continue
                    new_cell = ws.cell(row=r + 1, column=c + 1, value=value)

                    # Перенос форматов из ячейки Numbers
                    try:
                        if getattr(cell, "bold", False):
                            new_cell.font = Font(bold=True)
                        fill_color = getattr(cell, "fill_color", None)
                        if fill_color and getattr(fill_color, "rgb", None):
                            rgb = fill_color.rgb
                            if isinstance(rgb, tuple) and len(rgb) >= 3:
                                hex_color = "".join(f"{int(v):02X}" for v in rgb[:3])
                                new_cell.fill = PatternFill(
                                    start_color=hex_color,
                                    end_color=hex_color,
                                    fill_type="solid",
                                )
                        h_align = getattr(cell, "h_alignment", None)
                        if h_align:
                            align_map = {"left": "left", "center": "center",
                                         "right": "right"}
                            mapped = align_map.get(str(h_align).lower())
                            if mapped:
                                new_cell.alignment = Alignment(horizontal=mapped)
                    except Exception:
                        # Форматы — второстепенная информация, не прерываем конвертацию
                        pass

            # Ширины колонок
            for c in range(num_cols):
                try:
                    w = table.col_width(c)
                    if w:
                        # Примерный пересчёт пунктов в символы
                        ws.column_dimensions[get_column_letter(c + 1)].width = max(8, w / 7)
                except Exception:
                    pass

    stream = io.BytesIO()
    wb.save(stream)
    stream.seek(0)
    return stream


def sanitize_sheet_name(name: str) -> str:
    """Ограничивает имя листа согласно требованиям Excel."""
    if not name:
        return "Лист"
    for ch in "[]:*?/\\":
        name = name.replace(ch, "-")
    return name[:31]


if __name__ == "__main__":
    app.run(host="0.0.0.0", port=5000, debug=False)
